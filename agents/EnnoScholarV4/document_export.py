"""DOCX and PDF from an immutable saved manuscript, with consistent references."""
import html
import io
import re
from .models import public_link
from .scientific_visuals import figure_png


def blocks(text):
    lines=text.splitlines(); i=0
    while i<len(lines):
        line=lines[i]
        if not line.strip():
            i+=1; continue
        if i+1<len(lines) and '|' in line and re.fullmatch(r"[\s|:\-]+",lines[i+1]) and '-' in lines[i+1]:
            rows=[[c.strip() for c in line.strip().strip('|').split('|')]]; i+=2
            while i<len(lines) and '|' in lines[i] and lines[i].strip():
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')]); i+=1
            if all(len(r)==len(rows[0]) for r in rows):
                yield 'table',rows
            else:
                yield 'paragraph','\n'.join(' | '.join(r) for r in rows)
            continue
        heading=re.match(r'^(#{1,6})\s+(.+)',line)
        if heading:
            yield 'heading',(len(heading[1]),heading[2]); i+=1; continue
        if re.match(r'^\s*(?:[-*]|\d+\.)\s+',line):
            yield 'list',re.sub(r'^\s*(?:[-*]|\d+\.)\s+','',line); i+=1; continue
        paragraph=[line]; i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r'^#{1,6}\s|^\s*[-*]\s',lines[i]):
            if i+1<len(lines) and '|' in lines[i] and re.fullmatch(r'[\s|:\-]+',lines[i+1]):
                break
            paragraph.append(lines[i]); i+=1
        yield 'paragraph',' '.join(paragraph)


def plain_inline(text):
    text=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',text)
    return re.sub(r'\*\*([^*]+)\*\*|__([^_]+)__',lambda m:m[1] or m[2],text).replace('`','')


def export_document(state,version_id,fmt):
    revision=next((v for v in state.get('document_versions',[]) if v['id']==version_id),None)
    if revision is None:
        raise ValueError("Version enregistrée introuvable ; enregistrer le document avant l'export.")
    outline=revision['outline']
    sections=[(index, section) for index, section in enumerate(outline.get('sections',[]), 1) if section.get('drafts')]
    if not sections:
        raise ValueError("Aucune section rédigée et enregistrée dans cette version.")
    drafted_ids={section['id'] for _,section in sections}
    figures=[figure for figure in revision.get('figures',[]) if figure['section_id'] in drafted_ids]
    ids=list(dict.fromkeys([aid for _,s in sections for d in s.get('drafts',[])[-1:] for aid in d['article_ids']]+[
        state['evidence'][eid]['article_id'] for f in figures for eid in f['evidence_ids']]))
    for aid in ids:
        if aid not in state.get('articles',{}):
            raise ValueError("Référence de la version enregistrée absente du stockage.")
    # Same rendering as the workspace, from the drafts' raw identifiers: one heading per
    # section, grouped numbered citations, known metadata only.
    from .citation_render import render_manuscript, format_reference
    metadata={aid:{**state['articles'].get(aid,{}),**revision.get('article_metadata',{}).get(aid,{})} for aid in state.get('articles',{})}
    rendered=render_manuscript(state,[s for _,s in sections],metadata,state.get('evidence',{}))
    references={row['article_id']:row['number'] for row in rendered['references']}
    proof_numbers={row['evidence_id']:row['number'] for row in rendered['proofs']}
    for aid in ids:
        if aid not in references:
            references[aid]=len(references)+1
            rendered['references'].append({'number':references[aid],'article_id':aid,'text':format_reference(metadata[aid])})
    for figure in figures:
        for eid in figure['evidence_ids']:
            if eid not in proof_numbers:
                proof_numbers[eid]=f"P{len(proof_numbers)+1}"
                e=state['evidence'][eid]
                rendered['proofs'].append({'number':proof_numbers[eid],'evidence_id':eid,'reference':references.get(e['article_id']),
                    'section':e.get('section'),'page_number':e.get('page_number'),'quote':e.get('exact_quote','')[:500]})
    items=[('title',outline.get('title',"État de l'art scientifique"))]
    figure_number=0
    for num,section in sections:
        title=re.sub(r'^\s*\d+[.)]\s+','',section['title'])
        items.append(('heading',(1,f"{num} {title}")))
        body=rendered['sections'][section['id']]
        if re.search(r'\[(?:paper_|ev_|projet:)[\w-]+\]',body):
            raise ValueError("Citation non résolue dans la version exportée.")
        for kind,value in blocks(body):
            if kind=='heading':
                value=(min(4,value[0]+1),value[1])
            items.append((kind,value))
        for figure in figures:
            if figure['section_id']==section['id']:
                figure_number+=1; png=figure_png(figure)
                items.append(('image',(png,f'figure-{figure_number}.png',figure['title'])))
                caption=f"Figure {figure_number}. {figure['caption']}"
                if figure['origin']=='agent_synthesis':
                    caption+=' Schéma ou graphique de synthèse généré à partir des sources citées.'
                else:
                    caption+=f" Source : {figure['source_url']} ; {figure['source_version']} ; licence {figure['license_url']}."
                caption+=' Preuves : '+', '.join(proof_numbers[e] for e in figure['evidence_ids'])+'.'
                items.append(('caption',caption))
    items.append(('heading',(1,'Références bibliographiques')))
    for row in sorted(rendered['references'],key=lambda r:r['number']):
        items.append(('paragraph',f"[{row['number']}] {row['text']}"))
    if rendered['project_sources']:
        items.append(('heading',(1,'Sources du projet (EnnoDiagnostic)')))
        for row in rendered['project_sources']:
            location=' ; '.join(str(v) for v in (row.get('document_name'),row.get('section_title'),
                f"p. {row['page']}" if row.get('page') else None) if v)
            note=f" (localisateur du diagnostic : {row['stored_section_title']})" if row.get('stored_section_title') else ''
            items.append(('paragraph',f"{row['number']}. {location}{note}. Extrait : « {row['quote']} »"))
    items.append(('heading',(1,'Traçabilité des preuves')))
    for row in rendered['proofs']:
        e=state.get('evidence',{}).get(row['evidence_id'])
        if not e:
            raise ValueError('Preuve enregistrée absente du stockage.')
        location=e.get('section') or 'Section non identifiée'
        if e.get('page_number'):
            location+=f" ; page PDF {e['page_number']}"
        if e.get('printed_page'):
            location+=f" ; page imprimée {e['printed_page']}"
        quote=e['exact_quote'][:500]
        label='Extrait exact abrégé' if len(e['exact_quote'])>500 else 'Extrait exact'
        items.append(('paragraph',f"{row['number']}. Référence [{references[e['article_id']]}] ; {e.get('source_version','version inconnue')} ; {location} ; {e['source_url']}. {label} : « {quote} »"))
    return render_docx(items) if fmt=='docx' else render_pdf(items)


def render_docx(items):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document(); section=doc.sections[0]
    section.page_width=Inches(8.5); section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.8)
    section.left_margin=section.right_margin=Inches(.85)
    normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(8); normal.paragraph_format.line_spacing=1.2
    doc.styles['Title'].font.color.rgb=RGBColor(0,0,0)
    for level in range(1,5):
        doc.styles[f'Heading {level}'].font.color.rgb=RGBColor(0,0,0)
        doc.styles[f'Heading {level}'].paragraph_format.keep_with_next=True
    def paragraph(text,style=None):
        p=doc.add_paragraph(style=style)
        for token in re.split(r'(\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))',text):
            match=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',token)
            if match and public_link(match[2]):
                from docx.opc.constants import RELATIONSHIP_TYPE as RT
                link=OxmlElement('w:hyperlink'); link.set(qn('r:id'),p.part.relate_to(match[2],RT.HYPERLINK,is_external=True))
                run=OxmlElement('w:r'); value=OxmlElement('w:t'); value.text=match[1]; run.append(value); link.append(run); p._p.append(link)
            else:
                r=p.add_run(token[2:-2] if token.startswith('**') and token.endswith('**') else token); r.bold=token.startswith('**')
        return p
    for kind,value in items:
        if kind=='title':
            paragraph(value,'Title')
        elif kind=='heading':
            doc.add_heading(value[1],level=value[0])
        elif kind=='table':
            table=doc.add_table(rows=1,cols=len(value[0])); table.style='Light Shading Accent 1'
            repeat=OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for cell,text in zip(table.rows[0].cells,value[0]):
                cell.text=plain_inline(text)
            for row in value[1:]:
                for cell,text in zip(table.add_row().cells,row):
                    cell.text=plain_inline(text)
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(5); p.paragraph_format.space_before=Pt(5)
        elif kind=='image':
            import fitz
            pix=fitz.Pixmap(value[0]); ratio=pix.height/pix.width
            doc.add_picture(io.BytesIO(value[0]),width=Inches(min(6.5,8/ratio)))
            pic=doc.inline_shapes[-1]._inline.docPr; pic.set('descr',value[2])
        else:
            paragraph(value,'Caption' if kind=='caption' else 'List Bullet' if kind=='list' else None)
    footer=section.footer.paragraphs[0]; footer.alignment=2
    footer.add_run('EnnoScholar V4 | ')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    stream=io.BytesIO(); doc.save(stream); return stream.getvalue()


def render_pdf(items):
    import fitz
    archive=fitz.Archive(); fragments=[]
    def inline(value):
        value=html.escape(value)
        value=re.sub(r'\*\*([^*]+)\*\*',r'<b>\1</b>',value)
        return re.sub(r'\[([^\]]+)\]\((https?://[^)]+)\)',r'<a href="\2">\1</a>',value)
    for kind,value in items:
        if kind=='image':
            archive.add(value[0],value[1]); pix=fitz.Pixmap(value[0]); width=min(480,560*pix.width/pix.height)
            fragments.append(f'<p><img src="{value[1]}" width="{width}"/></p>')
        elif kind=='table':
            fragments.append('<table>'+''.join('<tr>'+''.join(f'<{"th" if i==0 else "td"}>{inline(c)}</{"th" if i==0 else "td"}>' for c in r)+'</tr>' for i,r in enumerate(value))+'</table>')
        elif kind=='heading':
            fragments.append(f'<h{value[0]+1}>{inline(value[1])}</h{value[0]+1}>')
        elif kind=='title':
            fragments.append('<h1>'+inline(value)+'</h1>')
        else:
            fragments.append(f'<p class="{kind}">'+inline(value)+'</p>')
    css='body {font-family:serif;font-size:11pt;line-height:1.3;} h1 {font-size:23pt;} h2 {font-size:16pt;} h3 {font-size:13pt;} p {margin-bottom:9pt;} table {border-collapse:collapse;width:100%;margin:10pt 0;} td,th {border-bottom:1px solid #aaa;padding:6pt;font-size:10pt;} th {background:#edf2f7;} .caption {font-size:10pt;} a {color:#24506d;}'
    story=fitz.Story(html='<html><body>'+''.join(fragments)+'</body></html>',user_css=css,archive=archive)
    stream=io.BytesIO(); writer=fitz.DocumentWriter(stream)
    story.write(writer,lambda rect_num,filled:(fitz.Rect(0,0,612,792),fitz.Rect(54,54,558,734),None))
    writer.close()
    with fitz.open(stream=stream.getvalue(),filetype='pdf') as doc:
        for i,page in enumerate(doc):
            page.insert_text((54,765),f'EnnoScholar V4 | {i+1} / {len(doc)}',fontsize=9)
        return doc.tobytes(garbage=4,deflate=True)
